import openpyxl
from openpyxl.styles import PatternFill, Font

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Student Template"

headers = ['RollNumber', 'FullName', 'Gender', 'MobileNo', 'DOB', 'FatherName', 'MotherName']
ws.append(headers)

yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
bold_font = Font(bold=True)

# First 3 columns are mandatory (RollNumber, FullName, Gender)
for col in range(1, 4):
    cell = ws.cell(row=1, column=col)
    cell.fill = yellow_fill
    cell.font = bold_font

# Next 4 are non-mandatory
for col in range(4, 8):
    cell = ws.cell(row=1, column=col)
    cell.font = bold_font

# Add a sample row
ws.append([101, 'John Doe', 'Male', '1234567890', '2010-05-15', 'Robert Doe', 'Jane Doe'])

wb.save("public/Student_Template.xlsx")
print("Successfully generated Student_Template.xlsx")
